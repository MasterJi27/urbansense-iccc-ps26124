from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
BG = RGBColor(244,246,248)
ACC = RGBColor(0x0B,0xB9,0x8A)
DARK = RGBColor(15,27,45)
MUTED = RGBColor(107,125,147)
WARN = RGBColor(183,121,15)
DANGER = RGBColor(229,72,77)

def bg(s):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG

def add_shape(slide, left, top, w, h, fill, line=None, radius=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    if line: shape.line.color.rgb = line; shape.line.width = Pt(1)
    if radius is not None:
        shape.adjustments[0] = 0.08
    return shape

def txt(shape, text, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT, italic=False):
    # guard: some calls passed PP_ALIGN as bold by mistake
    if isinstance(bold, type(PP_ALIGN.LEFT)):
        align = bold
        bold = False
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bool(bold)
    p.font.italic = bool(italic)
    p.alignment = align
    return p

def add_txt(slide, left, top, w, h, text, size=12, color=DARK, bold=False, align=PP_ALIGN.LEFT):
    tx = slide.shapes.add_textbox(left, top, w, h)
    txt(tx, text, size, color, bold, align)
    return tx

# SLIDE 1 Title
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
add_shape(s, Inches(0.6), Inches(0.6), Inches(12.1), Inches(1.1), RGBColor(255,255,255), RGBColor(231,236,241))
add_txt(s, Inches(0.9), Inches(0.7), Inches(7), Inches(0.4), "URBANSENSE", 28, DARK, True)
add_txt(s, Inches(0.9), Inches(1.05), Inches(7), Inches(0.4), "Structural Pulse — Continuous bridge-health screening using the buses already crossing the city.", 10, MUTED)
add_txt(s, Inches(0.9), Inches(1.4), Inches(7), Inches(0.2), "SIH 2026  •  PS 26124  BEL  •  Smart Automation  •  Theme: Mobile Urban Intelligence", 7, MUTED)
add_shape(s, Inches(9.2), Inches(0.7), Inches(3.2), Inches(0.9), ACC)
add_txt(s, Inches(9.3), Inches(0.8), Inches(3), Inches(0.25), "Buses as mobile structural sensors", 9, RGBColor(255,255,255), True, PP_ALIGN.CENTER)
add_txt(s, Inches(9.3), Inches(1.1), Inches(3), Inches(0.3), "Not collapse prediction — flag for inspection", 7, RGBColor(255,255,255), align=PP_ALIGN.CENTER)

# diagram
add_shape(s, Inches(0.6), Inches(2.1), Inches(12.1), Inches(4.7), RGBColor(255,255,255), RGBColor(231,236,241))
add_txt(s, Inches(0.8), Inches(2.3), Inches(11.7), Inches(0.3), "One bus → many intelligences", 11, DARK, True, PP_ALIGN.CENTER)
# bus center
add_shape(s, Inches(5.9), Inches(2.8), Inches(1.5), Inches(0.5), DARK)
add_txt(s, Inches(5.9), Inches(2.85), Inches(1.5), Inches(0.4), "PUBLIC BUS", 8, RGBColor(255,255,255), True, PP_ALIGN.CENTER)
# arrows down
for x in [Inches(3.5), Inches(6.65), Inches(9.5)]:
    add_txt(s, x, Inches(3.6), Inches(1), Inches(0.2), "↓", 14, MUTED, align=PP_ALIGN.CENTER)
add_shape(s, Inches(2.5), Inches(3.9), Inches(2), Inches(0.7), RGBColor(248,250,251), RGBColor(231,236,241))
add_txt(s, Inches(2.6), Inches(4.0), Inches(1.8), Inches(0.2), "Camera", 9, DARK, True, PP_ALIGN.CENTER)
add_txt(s, Inches(2.6), Inches(4.25), Inches(1.8), Inches(0.2), "Road defects", 7, MUTED, PP_ALIGN.CENTER)
add_shape(s, Inches(5.65), Inches(3.9), Inches(2), Inches(0.7), RGBColor(248,250,251), RGBColor(231,236,241))
add_txt(s, Inches(5.75), Inches(4.0), Inches(1.8), Inches(0.2), "GPS", 9, DARK, True, PP_ALIGN.CENTER)
add_txt(s, Inches(5.75), Inches(4.25), Inches(1.8), Inches(0.2), "Location", 7, MUTED, PP_ALIGN.CENTER)
add_shape(s, Inches(8.8), Inches(3.9), Inches(2), Inches(0.7), RGBColor(248,250,251), RGBColor(231,236,241))
add_txt(s, Inches(8.9), Inches(4.0), Inches(1.8), Inches(0.2), "IMU", 9, DARK, True, PP_ALIGN.CENTER)
add_txt(s, Inches(8.9), Inches(4.25), Inches(1.8), Inches(0.2), "Vibration", 7, MUTED, PP_ALIGN.CENTER)
add_txt(s, Inches(5.9), Inches(4.8), Inches(1.5), Inches(0.2), "↓", 14, MUTED, PP_ALIGN.CENTER)
add_shape(s, Inches(5), Inches(5.1), Inches(3.3), Inches(0.7), ACC)
add_txt(s, Inches(5.1), Inches(5.15), Inches(3.1), Inches(0.25), "URBANSENSE  →  Fusion (spatial + source diversity)", 8, RGBColor(255,255,255), True, PP_ALIGN.CENTER)
add_shape(s, Inches(2.0), Inches(6.0), Inches(3.5), Inches(0.6), RGBColor(255,255,255), RGBColor(231,236,241))
add_txt(s, Inches(2.1), Inches(6.05), Inches(3.3), Inches(0.5), "Road condition\nPotholes, waterlogging, missing assets", 7, DARK, align=PP_ALIGN.CENTER)
add_shape(s, Inches(7.8), Inches(6.0), Inches(3.5), Inches(0.6), RGBColor(232,247,241), RGBColor(11,185,138))
add_txt(s, Inches(7.9), Inches(6.05), Inches(3.3), Inches(0.5), "Structure health ★\nBridge / ROB / Flyover screening", 7, DARK, True, PP_ALIGN.CENTER)

# SLIDE 2 Problem
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
add_txt(s, Inches(0.6), Inches(0.4), Inches(7), Inches(0.3), "Problem → why buses?", 11, MUTED)
add_txt(s, Inches(0.6), Inches(0.7), Inches(7), Inches(0.5), "India has a huge inventory of bridges/ROBs/flyovers.\nFixed instrumentation on every structure is expensive.", 14, DARK, True)
add_shape(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5), RGBColor(255,255,255), RGBColor(231,236,241))
add_txt(s, Inches(0.8), Inches(1.8), Inches(5.5), Inches(0.25), "Today (others):", 9, MUTED, True)
add_txt(s, Inches(0.8), Inches(2.1), Inches(5.5), Inches(0.8), "bus camera → pothole → dashboard\n(camera carrier only)", 10, DARK)
add_txt(s, Inches(0.8), Inches(3.1), Inches(5.5), Inches(0.4), "Cost: 10–20L per bridge for fixed SHM", 8, DANGER, True)
add_txt(s, Inches(0.8), Inches(3.7), Inches(5.5), Inches(1.5), "Gap: Periodic inspection → visible deterioration tak pata hi nahi chalta.\nResearch: vehicle dynamics / road-profile / sensor noise are major issues — one pass ≠ diagnosis.", 7, MUTED)
add_shape(s, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5), ACC)
add_txt(s, Inches(7.0), Inches(1.8), Inches(5.5), Inches(0.25), "Our jump:", 9, RGBColor(255,255,255), True)
add_txt(s, Inches(7.0), Inches(2.1), Inches(5.5), Inches(0.8), "What if the bus can also sense the health\nof the infrastructure it physically travels over?", 10, RGBColor(255,255,255), True)
add_txt(s, Inches(7.0), Inches(3.1), Inches(5.5), Inches(0.4), "Fleet already traverses structures → temporary sensing nodes", 8, RGBColor(255,255,255))
add_txt(s, Inches(7.0), Inches(3.8), Inches(5.5), Inches(1.5), "Benefit: Network-level screening without installing on every structure.\nBEL relevant: electronics + sensor signal processing.", 7, RGBColor(255,255,255))

# SLIDE 3 Sense Learn Flag
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
add_txt(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.3), "Structural Pulse — Sense  •  Learn  •  Flag  (not collapse prediction)", 11, MUTED)
add_txt(s, Inches(0.6), Inches(0.7), Inches(12), Inches(0.4), "Continuous screening → inspection recommendation", 16, DARK, True)
for i, (title, body, col) in enumerate([
    ("SENSE", "Smartphone IMU (50Hz) + GPS\nrecords bridge-crossing dynamics\n70m geofence, 1–3s batch", ACC),
    ("LEARN", "Per-structure EMA baseline\nf_dom + RMS + crest\nLongitudinal history 80 pts", DARK),
    ("FLAG", "Persistent drift across buses\n→ ANOMALOUS\nAuto WO-BR-* + inspection", WARN),
]):
    x = Inches(0.6 + i*4.3)
    add_shape(s, x, Inches(1.4), Inches(3.9), Inches(4.5), RGBColor(255,255,255), RGBColor(231,236,241))
    add_shape(s, x+Inches(0.2), Inches(1.6), Inches(3.5), Inches(0.5), col)
    add_txt(s, x+Inches(0.25), Inches(1.65), Inches(3.4), Inches(0.4), title, 10, RGBColor(255,255,255), True, PP_ALIGN.CENTER)
    add_txt(s, x+Inches(0.3), Inches(2.3), Inches(3.3), Inches(2.5), body, 9, DARK, align=PP_ALIGN.CENTER)
add_txt(s, Inches(0.6), Inches(6.2), Inches(12.1), Inches(0.5), "Correct claim:  “This bridge's observed dynamic response has changed enough from its baseline to warrant inspection.”   •   Screening layer, not certified inspection.", 7, MUTED, align=PP_ALIGN.CENTER)

# SLIDE 4 Architecture
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
add_txt(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.3), "Architecture — one observation stream, not a separate product", 11, MUTED)
add_shape(s, Inches(0.6), Inches(0.9), Inches(12.1), Inches(5.8), RGBColor(255,255,255), RGBColor(231,236,241))
add_txt(s, Inches(0.8), Inches(1.0), Inches(11.7), Inches(0.2), "BUS", 10, DARK, True, PP_ALIGN.CENTER)
add_txt(s, Inches(0.8), Inches(1.4), Inches(3.5), Inches(0.7), "CAMERA\nRoad AI → pothole", 8, DARK, align=PP_ALIGN.CENTER)
add_txt(s, Inches(4.9), Inches(1.4), Inches(3.5), Inches(0.7), "GPS\nLocation + speed + direction", 8, DARK, align=PP_ALIGN.CENTER)
add_txt(s, Inches(9.0), Inches(1.4), Inches(3.5), Inches(1.0), "IMU\nVibration → RMS/peak/crest\n+ FFT f_dom", 8, DARK, align=PP_ALIGN.CENTER)
add_txt(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(0.2), "└───────────────┼───────────────┘", 10, MUTED, PP_ALIGN.CENTER)
add_txt(s, Inches(0.8), Inches(2.7), Inches(11.7), Inches(0.3), "↓  OBSERVATION (phone accel + GPS, ~1KB JSON, ~1.6s on bridge)", 8, ACC, True, PP_ALIGN.CENTER)
add_shape(s, Inches(4.5), Inches(3.2), Inches(4.3), Inches(0.6), DARK)
add_txt(s, Inches(4.6), Inches(3.3), Inches(4.1), Inches(0.4), "URBAN FUSION  (spatial + temporal + source diversity)", 8, RGBColor(255,255,255), True, PP_ALIGN.CENTER)
add_txt(s, Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.2), "↓", 14, MUTED, PP_ALIGN.CENTER)
add_txt(s, Inches(0.8), Inches(4.3), Inches(11.7), Inches(0.3), "STRUCTURE / ROAD STATE  →  Health score  •  Work Order  •  GIS map", 9, DARK, True, PP_ALIGN.CENTER)
add_txt(s, Inches(0.8), Inches(4.9), Inches(11.7), Inches(0.4), "One bus tells:  road deteriorating  +  bridge anomalous  +  traffic rising  +  pedestrian high", 7, MUTED, align=PP_ALIGN.CENTER)
add_txt(s, Inches(0.8), Inches(5.5), Inches(11.7), Inches(0.6), "Per-structure asset:  BRIDGE-DEL-042  •  dynamic baseline (f_dom, RMS, spectral)  •  environmental (speed, direction, time)  •  state NORMAL→WATCH→ANOMALOUS→INSPECTION REQUIRED", 7, MUTED, align=PP_ALIGN.CENTER)
add_txt(s, Inches(0.6), Inches(6.9), Inches(12.1), Inches(0.3), "Reuse: Observation → Fusion → UrbanEvent • No new platform.   DOPT safe: no face/plate, only vibration.", 7, MUTED, align=PP_ALIGN.CENTER)

# SLIDE 5 Demo loop
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
add_txt(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.3), "Demo loop — 60 days in 60 seconds (killer version)", 11, MUTED)
for i, (day, bus, res) in enumerate([("DAY 1","BUS-017","Baseline created"),("DAY 14","BUS-042","Matches baseline"),("DAY 60","BUS-017/042/103","Statistically significant drift\n⚠ STRUCTURAL ANOMALY → Inspection")] ):
    x = Inches(0.6 + i*4.3)
    add_shape(s, x, Inches(0.9), Inches(3.9), Inches(4.2), RGBColor(255,255,255), RGBColor(231,236,241))
    add_txt(s, x+Inches(0.2), Inches(1.0), Inches(3.5), Inches(0.25), day, 9, MUTED, True, PP_ALIGN.CENTER)
    add_txt(s, x+Inches(0.2), Inches(1.3), Inches(3.5), Inches(0.3), bus, 10, DARK, True, PP_ALIGN.CENTER)
    add_txt(s, x+Inches(0.2), Inches(1.6), Inches(3.5), Inches(0.2), "Bridge X", 8, MUTED, PP_ALIGN.CENTER)
    add_txt(s, x+Inches(0.2), Inches(1.95), Inches(3.5), Inches(0.2), "↓  Acceleration signature  ↓", 7, ACC, align=PP_ALIGN.CENTER)
    add_txt(s, x+Inches(0.2), Inches(2.3), Inches(3.5), Inches(0.4), "Baseline" if i==0 else ("Matches" if i==1 else "Repeated measurements"), 9, DARK, True, PP_ALIGN.CENTER)
    add_shape(s, x+Inches(0.6), Inches(2.9), Inches(2.7), Inches(0.6), ACC if i<2 else DANGER)
    add_txt(s, x+Inches(0.65), Inches(3.0), Inches(2.6), Inches(0.4), res, 7, RGBColor(255,255,255), True, PP_ALIGN.CENTER)
add_txt(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(0.3), "The word DRIFT is important — we never claim ‘will collapse’.", 9, DANGER, True, PP_ALIGN.CENTER)
add_txt(s, Inches(0.6), Inches(6.0), Inches(12.1), Inches(0.8), "Judge Q: ‘Why buses vs sensors on bridge?’ → ‘Not another network to maintain. Fleet already crosses daily — network-scale without installing on every structure.’\nJudge Q: ‘Bridge vs bus?’ → Control for direction/speed + persistent across buses.", 7, MUTED)

# SLIDE 6 Innovation scores
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
add_txt(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.3), "Why this wins — honest scores", 11, MUTED)
scores = [("Impact",9.5),("Fit with PS",9.5),("Technical depth",9),("BEL relevance",9),("Demo potential",9),("System novelty",8.5)]
for i,(k,v) in enumerate(scores):
    y = Inches(1.0 + i*0.85)
    add_txt(s, Inches(0.8), y, Inches(3), Inches(0.25), k, 10, DARK, True)
    add_shape(s, Inches(4.0), y+Inches(0.05), Inches(6), Inches(0.2), RGBColor(231,236,241))
    add_shape(s, Inches(4.0), y+Inches(0.05), Inches(6*v/10), Inches(0.2), ACC)
    add_txt(s, Inches(10.3), y, Inches(1), Inches(0.25), f"{v}/10", 10, DARK, True)
add_txt(s, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.5), "Science novelty 5/10 — drive-by SHM exists (2023 IoT, 2026 smartphone modal). Our novelty = fleet-scale screening as one layer of urban intelligence.", 7, MUTED, align=PP_ALIGN.CENTER)
add_txt(s, Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.3), "Warning: Build screening demonstrator (2–3 phones same flyover → f_dom/RMS → GIS yellow/red), not collapse proof.", 8, DANGER, True, PP_ALIGN.CENTER)

prs.save("../docs/URBANSENSE_Structural_Pulse_Deck.pptx")
print("saved")
